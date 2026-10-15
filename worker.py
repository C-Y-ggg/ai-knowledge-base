import redis
import time
import sys
import pymupdf
import os
from openai import OpenAI
from app.database import SessionLocal
from app.models import Knowledge, Chunk
from app.llm_utils import get_embedding
from docx import Document
import pandas as pd


def extract_text_from_file(file_path: str) -> str:
    """支持 PDF / Word / Excel / 纯文本的多格式解析引擎"""
    if not os.path.exists(file_path):
        return ""

    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    try:
        # 1. 处理 PDF
        if ext == '.pdf':
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
            doc.close()

        # 2. 处理 Word (docx)
        elif ext == '.docx':
            doc = Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
            # 👇 进阶：Word 里的表格也别放过！
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells])
                    text += row_text + "\n"

        # 3. 处理 Excel (xlsx / xls)
        elif ext in ['.xlsx', '.xls']:
            # 读取所有 Sheet
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                # 把表格转换成 Markdown 格式的文本，大模型最喜欢这种格式
                text += f"\n【工作表: {sheet_name}】\n"
                text += df.to_markdown(index=False) + "\n"

        # 4. 处理纯文本 (txt / md)
        elif ext in ['.txt', '.md']:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

        else:
            return "暂不支持该文件格式"

        return text

    except Exception as e:
        print(f"❌ 解析文件 {file_path} 时出错: {e}")
        return ""

def split_text_into_chunks(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    """改进版：按段落优先切分，避免切断表格和标题"""
    # 先按双换行（段落）切分
    paragraphs = text.split("\n\n")
    chunks = []
    current_chunk = ""

    for para in paragraphs:
        # 如果单段落本身就超过 chunk_size，直接切
        if len(para) > chunk_size:
            if current_chunk:
                chunks.append(current_chunk.strip())
                current_chunk = ""
            # 长段落内部按 chunk_size 切
            for i in range(0, len(para), chunk_size - overlap):
                chunks.append(para[i:i + chunk_size].strip())
        # 如果加上这一段会超长，先保存当前块
        elif len(current_chunk) + len(para) > chunk_size:
            chunks.append(current_chunk.strip())
            current_chunk = para
        else:
            current_chunk += "\n\n" + para

    if current_chunk.strip():
        chunks.append(current_chunk.strip())

    return [c for c in chunks if c]


# ================= 后厨主循环 =================
r = redis.Redis(host='localhost', port=6379, decode_responses=True, socket_timeout=10)
print("👨‍🍳 后台 Worker 已启动，正在监听 Redis 队列...")

while True:
    try:
        task = r.blpop("knowledge_queue", timeout=5)
        if task:
            queue_name, knowledge_id = task
            print(f"\n🔔 收到任务！knowledge_id: {knowledge_id}")

            db = SessionLocal()
            try:
                knowledge = db.query(Knowledge).filter(Knowledge.id == int(knowledge_id)).first()
                if not knowledge:
                    continue

                knowledge.parse_status = "processing"
                db.commit()
                print(f"⏳ 开始解析文件: {knowledge.file_name}")

                raw_text = extract_text_from_file(knowledge.file_path)
                if not raw_text:
                    raise Exception("文档内容为空或无法解析")

                chunks = split_text_into_chunks(raw_text, chunk_size=500, overlap=50)
                print(f"📄 文件解析完成，共切分出 {len(chunks)} 个文本块")

                # 清空旧数据
                db.query(Chunk).filter(Chunk.knowledge_id == knowledge.id).delete()

                # 逐块呼叫翻译官，生成向量并入库
                success_count = 0  # 新增：成功计数器
                for i, chunk_content in enumerate(chunks):
                    print(f"  🧠 正在翻译第 {i + 1}/{len(chunks)} 块...")
                    vector = get_embedding(chunk_content)

                    if not vector:
                        print("  ⚠️ 向量生成失败，跳过该块")
                        continue

                    new_chunk = Chunk(
                        knowledge_id=knowledge.id,
                        content=chunk_content,
                        embedding=vector
                    )
                    db.add(new_chunk)
                    success_count += 1  # 新增：成功则加 1
                    time.sleep(0.5)

                # 新增：如果一块都没成功，就抛出异常，让外层捕获并把状态改为 failed
                if success_count == 0 and len(chunks) > 0:
                    raise Exception("所有文本块的向量化均失败，请检查 API Key 或余额")

                knowledge.parse_status = "completed"
                db.commit()
                print(f"✅ 处理完成！knowledge_id: {knowledge_id} 已向量化并入库，成功生成 {success_count} 个向量")

            except Exception as e:
                print(f"❌ 处理任务时出错: {e}")
                db.rollback()
                knowledge.parse_status = "failed"
                db.commit()
            finally:
                db.close()

    except redis.exceptions.TimeoutError:
        pass
    except KeyboardInterrupt:
        print("\n👋 Worker 已停止。")
        sys.exit(0)
    except Exception as e:
        print(f"❌ Redis 连接出错: {e}")
        time.sleep(5)